"""Inspect ZeoDiff's public release without importing its application code.

The audit reads workbooks and checkpoint metadata. The optional replay reads
the audit's numeric arrays, restricts deserialization to known numeric types,
and evaluates the published feedforward network on the full dataset. Replay
metrics are diagnostics, not an independent-test reproduction.
"""
from __future__ import annotations

import argparse
from collections import OrderedDict
import io
import json
from pathlib import Path
import pickle
import zipfile

import numpy as np


FEATURES = ["FDSi", "PLD", "PLD/LCD", "Vacc", "Tort", "AvgA", "StdA", "MaxA", "ASA"]


def write_json(path: Path, data: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def read_checkpoint(path: Path) -> dict:
    """Decode the author's float32 tensor archive using a small allowlist."""
    with zipfile.ZipFile(path) as archive:
        prefix = next(name[:-8] for name in archive.namelist() if name.endswith("data.pkl"))
        if archive.read(prefix + "byteorder").strip() != b"little":
            raise ValueError("Only little-endian author tensors are supported")

        def rebuild(storage, offset, size, stride, requires_grad, hooks):
            size, stride = tuple(size), tuple(stride)
            # Only contiguous tensors from this release are needed.
            expected = 1
            for dim, step in reversed(list(zip(size, stride))):
                if step != expected:
                    raise ValueError("Non-contiguous tensor in author checkpoint")
                expected *= dim
            return storage[offset:offset + expected].reshape(size).copy()

        class NumericUnpickler(pickle.Unpickler):
            def find_class(self, module, name):
                allowed = {
                    ("collections", "OrderedDict"): OrderedDict,
                    ("torch._utils", "_rebuild_tensor_v2"): rebuild,
                    ("torch", "FloatStorage"): "float32",
                }
                if (module, name) not in allowed:
                    raise ValueError(f"Unsupported checkpoint type: {module}.{name}")
                return allowed[module, name]

            def persistent_load(self, value):
                kind, dtype, key, device, length = value
                if kind != "storage" or dtype != "float32" or device != "cpu":
                    raise ValueError("Unsupported checkpoint storage")
                data = np.frombuffer(archive.read(prefix + "data/" + key), dtype="<f4")
                if len(data) != length:
                    raise ValueError("Storage length mismatch")
                return data

        return NumericUnpickler(io.BytesIO(archive.read(prefix + "data.pkl"))).load()


def metrics(y: np.ndarray, prediction: np.ndarray) -> dict:
    residual = y - prediction
    return {
        "n": len(y),
        "mae": float(np.mean(np.abs(residual))),
        "rmse": float(np.sqrt(np.mean(residual ** 2))),
        "r2": float(1 - np.sum(residual ** 2) / np.sum((y - np.mean(y)) ** 2)),
    }


def audit(root: Path, output: Path) -> None:
    from openpyxl import load_workbook

    release = root / "author_release"
    dataset = release / "Dataset/zeolite_diffusion_dataset.xlsx"
    workbook = load_workbook(dataset, read_only=True, data_only=True)
    sheet = workbook["Sheet1"]
    rows = list(sheet.iter_rows(values_only=True))
    columns = list(rows[0])
    records = [dict(zip(columns, row)) for row in rows[1:] if row[0] is not None]
    workbook.close()
    ids = np.array([str(record["Zeolites"]).strip().lower() for record in records])
    x = np.array([[record[name] for name in FEATURES] for record in records], dtype=float)
    ds = np.array([record["Ds"] for record in records], dtype=float)
    if not np.isfinite(x).all() or not np.isfinite(ds).all() or (ds <= 0).any():
        raise ValueError("Incomplete features or non-positive diffusion labels")
    np.savez(root / "audit_arrays.npz", ids=ids, features=x, ds_scaled=ds)

    source = load_workbook(root / "source_data.xlsx", read_only=True, data_only=True)
    source_sheets = [{"name": s.title, "rows": s.max_row, "columns": s.max_column}
                     for s in source.worksheets]
    plot_metrics = {}
    for name, start, pairs in [
        ("Supplementary data 3", 4, [("initial_ann_train", 0), ("initial_ann_test", 2)]),
        ("Supplementary data 10", 4, [("final_ann_train", 0), ("final_ann_test", 2)]),
        ("Supplementary data 15", 5, [("xgb_train", 0), ("xgb_test", 2),
                                       ("lgbm_train", 5), ("lgbm_test", 7),
                                       ("rf_train", 10), ("rf_test", 12)]),
    ]:
        sheet_rows = list(source[name].iter_rows(min_row=start, values_only=True))
        for label, index in pairs:
            pairs_array = np.array([[row[index], row[index + 1]] for row in sheet_rows
                                    if isinstance(row[index], (int, float))
                                    and isinstance(row[index + 1], (int, float))], dtype=float)
            plot_metrics[label] = metrics(pairs_array[:, 1], pairs_array[:, 0])
            np.save(root / f"source_{label}.npy", pairs_array)
    table9 = list(source["Supplementary data 9"].iter_rows(min_row=3, values_only=True))
    table9 = [{"model": r[0], "r2": r[1], "rmse": r[2]} for r in table9 if r[0]]
    shap_ids = [str(row[0]).strip().lower() for row in
                source["Supplementary data 11"].iter_rows(min_row=3, values_only=True)
                if row[0] is not None]
    np.save(root / "author_training_ids.npy", np.asarray(shap_ids))
    source.close()
    checkpoint = read_checkpoint(release / "ZeoDiff_HTP_master/HTP_scripts/models/ann_final_model.pth")
    cp_features = checkpoint["feature_names"]
    if cp_features != FEATURES:
        raise ValueError("Author checkpoint feature order changed")
    data = {
        "checked_on": "2026-09-29",
        "paper": {"doi": "10.1038/s41467-026-71698-0",
                  "url": "https://www.nature.com/articles/s41467-026-71698-0"},
        "release": {"zenodo": "https://zenodo.org/records/18858951", "version": "v1.0",
                    "repository": "https://github.com/WangXiaobao-j/ZeoDiff_HTP",
                    "downloads": json.loads((root / "download_manifest.json").read_text(encoding="utf-8"))},
        "dataset": {
            "rows": len(records), "unique_ids": len(set(ids)), "source_columns": columns,
            "native_d0": FEATURES, "initial_d0": ["PLD", "FDSi", "Vacc", "ASA"],
            "all_features_finite": True, "all_targets_positive": True,
            "target": "Ds is D / (1e-8 m^2/s); plot sheets use log10(Ds)",
            "feature_ranges": {name: {"min": float(x[:, i].min()), "max": float(x[:, i].max())}
                               for i, name in enumerate(FEATURES)},
            "target_range_scaled": {"min": float(ds.min()), "max": float(ds.max())},
            "id_sources": {"alphabetic_ids": int(sum(s.isalpha() for s in ids)),
                           "numeric_ids": int(sum(s.isdigit() for s in ids))},
        },
        "source_data": {"sheets": source_sheets, "reported_table9": table9,
                        "metrics_recomputed_from_plot_pairs": plot_metrics,
                        "shap_ids": {"n": len(shap_ids), "unique": len(set(shap_ids)),
                                     "not_in_main_dataset": sorted(set(shap_ids) - set(ids))}},
        "checkpoint": {"feature_names": cp_features, "hyperparams": checkpoint["hyperparams"],
                       "optimizer": {k: v for k, v in checkpoint["optimizer_state_dict"]["param_groups"][0].items()
                                     if k not in {"params"}},
                       "weight_shapes": {k: list(v.shape) for k, v in checkpoint["model_state_dict"].items()}},
        "status": "source_audit_complete; original split and training recipe need reconciliation",
        "discovery_jobs_submitted": False,
    }
    write_json(output, data)
    print(json.dumps({"rows": len(records), "native_d0_count": len(FEATURES),
                      "published_plot_metrics": plot_metrics, "output": str(output)}, ensure_ascii=True))


def replay(root: Path, output: Path) -> None:
    # This operation uses numeric audit arrays; it does not read or edit workbooks.
    from joblib.numpy_pickle import NumpyUnpickler, NumpyArrayWrapper
    from sklearn.preprocessing import StandardScaler
    from sklearn.model_selection import train_test_split

    class RestrictedScalerUnpickler(NumpyUnpickler):
        def find_class(self, module, name):
            allowed = {
                ("sklearn.preprocessing._data", "StandardScaler"): StandardScaler,
                ("joblib.numpy_pickle", "NumpyArrayWrapper"): NumpyArrayWrapper,
                ("numpy", "ndarray"): np.ndarray,
                ("numpy", "dtype"): np.dtype,
                ("numpy._core.multiarray", "scalar"): np.core.multiarray.scalar,
                ("numpy.core.multiarray", "scalar"): np.core.multiarray.scalar,
            }
            if (module, name) not in allowed:
                raise ValueError(f"Unsupported scaler type: {module}.{name}")
            return allowed[module, name]

    model_dir = root / "author_release/ZeoDiff_HTP_master/HTP_scripts/models"
    with (model_dir / "scaler.pkl").open("rb") as handle:
        scaler = RestrictedScalerUnpickler(str(model_dir / "scaler.pkl"), handle, mmap_mode=None).load()
    checkpoint = read_checkpoint(model_dir / "ann_final_model.pth")
    with np.load(root / "audit_arrays.npz", allow_pickle=False) as arrays:
        ids, x, ds = arrays["ids"], arrays["features"], arrays["ds_scaled"]
    weights = checkpoint["model_state_dict"]
    prediction = scaler.transform(x).astype(np.float32)
    for layer in ["hidden1", "hidden2", "hidden3", "output"]:
        prediction = prediction @ weights[layer + ".weight"].T + weights[layer + ".bias"]
        if layer != "output":
            prediction = np.maximum(prediction, 0)
    prediction = prediction[:, 0].astype(float)
    report = json.loads(output.read_text(encoding="utf-8"))
    report["checkpoint_replay_diagnostic"] = {
        "evaluation_scope": "full dataset, mixes training and test; not independent performance",
        "implementation": "NumPy forward pass of published weights; no application code executed",
        "scaler_n_samples_seen": int(scaler.n_samples_seen_),
        "scaler_mean": scaler.mean_.tolist(), "scaler_scale": scaler.scale_.tolist(),
        "as_ln_target": metrics(np.log(ds), prediction),
        "as_log10_target": metrics(np.log10(ds), prediction),
        "author_gui_inverse": "exp(output) * 1e-8 m^2/s",
    }
    author_ids = np.load(root / "author_training_ids.npy", allow_pickle=False)
    train_mask = np.isin(ids, author_ids)
    seed_train, seed_test = train_test_split(np.arange(len(ids)), test_size=0.1, random_state=3407)
    mean_difference = float(np.max(np.abs(x[train_mask].mean(axis=0) - scaler.mean_)))
    scale_difference = float(np.max(np.abs(x[train_mask].std(axis=0) - scaler.scale_)))
    seed_matches = set(ids[seed_train]) == set(author_ids)
    split_matches = (len(set(author_ids)) == int(scaler.n_samples_seen_)
                     and int(train_mask.sum()) == len(author_ids)
                     and mean_difference < 1e-10 and scale_difference < 1e-10)
    report["native_split_reconstruction"] = {
        "method": "SHAP training-table IDs only, matched to author scaler mean and population standard deviation",
        "shap_values_used_as_features": False,
        "train_count": int(train_mask.sum()), "test_count": int((~train_mask).sum()),
        "scaler_mean_max_abs_difference": mean_difference,
        "scaler_scale_max_abs_difference": scale_difference,
        "seed_3407_same_training_ids": seed_matches,
        "membership_verified": split_matches,
        "original_split_generation_and_order_verified": False,
        "checkpoint_metrics_log10_ds": {
            name: metrics(np.log10(ds[mask]), prediction[mask] / np.log(10))
            for name, mask in [("train", train_mask), ("test", ~train_mask)]
        },
    }
    if split_matches:
        write_json(root / "reconstructed_author_split.json", {
            "method": report["native_split_reconstruction"]["method"],
            "training_ids": ids[train_mask].tolist(), "test_ids": ids[~train_mask].tolist(),
        })
    report["status"] = "native descriptors and checkpoint replay audited; training recipe and plot-data version unresolved"
    np.save(root / "checkpoint_output.npy", prediction)
    write_json(output, report)
    print(json.dumps({"replay": report["checkpoint_replay_diagnostic"],
                      "native_split": report["native_split_reconstruction"]}, ensure_ascii=True))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("operation", choices=["audit", "replay"])
    parser.add_argument("--raw", type=Path, default=Path("research/datasets/raw/zeodiff-2026"))
    parser.add_argument("--output", type=Path, default=Path("research/reports/zeodiff_source_audit_20260929.json"))
    args = parser.parse_args()
    (audit if args.operation == "audit" else replay)(args.raw, args.output)
